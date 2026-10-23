from __future__ import annotations

import importlib.util
import io
import logging
import threading
import time
import warnings

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from PIL import Image, ImageOps, UnidentifiedImageError

from .analysis import analyze

app = FastAPI(title="Maíz Raíz", version="0.1.0")
from .network_info import router as network_router
app.include_router(network_router)
lock = threading.Lock()
MAX_BYTES = 25 * 1024 * 1024
Image.MAX_IMAGE_PIXELS = 85_000_000


@app.post("/api/yield/experiment")
def yield_experiment(mode: str = Form("imagen"), kernel: str = Form("rbf"), c: float = Form(1),
                     gamma: str = Form("scale"), degree: int = Form(3), threshold: str = Form(""),
                     min_grains: int = Form(25), excluded: str = Form(""), nested: bool = Form(False)):
    import json, uuid
    from datetime import datetime, timezone
    from .yield_experiment import MODES
    from .svm_experiment import run
    from .model_registry import save_candidate
    if mode not in MODES:
        raise HTTPException(422, "Modo no reconocido.")
    with history_db() as db:
        rows = db.execute("SELECT id, payload FROM samples WHERE deleted = 0 ORDER BY id").fetchall()
    if not lock.acquire(blocking=False):
        raise HTTPException(429, "Hay un análisis o experimento en curso.")
    try:
        report = run([(identifier, json.loads(payload)) for identifier, payload in rows], mode,
                     kernel=kernel, c=c, gamma=gamma, degree=degree, threshold=float(threshold) if threshold.strip() else None,
                     min_grains=min_grains, excluded=excluded.split(",") if excluded.strip() else (), nested=nested,
                     artifact_sink=lambda fitted, meta: save_candidate(HISTORY_PATH.parent, fitted, meta))
    except ValueError as exc:
        raise HTTPException(422, str(exc))
    finally:
        lock.release()
    report["created"] = datetime.now(timezone.utc).isoformat()
    if report["status"] == "completed":
        report["id"] = str(uuid.uuid4())
        directory = HISTORY_PATH.parent / "yield_experiments"
        directory.mkdir(parents=True, exist_ok=True)
        current_ids = {r["sample_id"] for r in report["samples"]}
        prior = []
        for previous in directory.glob("*.json"):
            try:
                saved = json.loads(previous.read_text(encoding="utf-8"))
                if current_ids.intersection(saved.get("test_ids", [])):
                    prior.append(saved.get("id", previous.stem))
            except (ValueError, OSError):
                continue
        report["previous_test_exposures"] = len(prior)
        report["previous_reports"] = prior
        destination = directory / (report["id"] + ".json")
        temporary = destination.with_suffix(".tmp")
        temporary.write_text(json.dumps(report, ensure_ascii=False, allow_nan=False), encoding="utf-8")
        temporary.replace(destination)
    return report


@app.get('/api/models')
def model_versions():
    from .model_registry import models
    return models(HISTORY_PATH.parent)


@app.post('/api/models/{identifier}/activate')
def activate_model(identifier: str):
    from .model_registry import activate
    try:
        return activate(HISTORY_PATH.parent, identifier)
    except ValueError as exc:
        raise HTTPException(422, str(exc))


def result_summary(payload):
    from .weight import estimate, measurement
    from .model_registry import classify
    from .yield_experiment import family, number
    payload = dict(payload)
    result = payload.get('result') or {}
    if not payload.get('metadata') and payload.get('metadata_selection') != 'explicit' and METADATA_PATH.exists():
        saved = __import__('json').loads(METADATA_PATH.read_text(encoding='utf-8'))
        key = next((c for c in saved['columns'] if c in ['nombre_imagen','nombre_nuevo','familia','muestra']), saved['columns'][0])
        matched = [r for r in saved['rows'] if family(r.get(key,'')) == family(result.get('sample',''))]
        if len(matched) == 1:
            payload['metadata'] = {**saved, 'rows':matched, 'key':key, 'identifier':matched[0][key]}
    grains = result.get('grains') or []
    inputs = payload.get('weight_inputs') or {'mode': 'auto'}
    record = dict(payload)
    signature = [[g.get(k) for k in ['grano', 'largo', 'ancho', 'area']] for g in grains]
    exact = inputs.get('measured') and inputs.get('grain_signature') == signature and inputs.get('mode') == 'lot'
    measured = None
    if exact:
        record['weight_measurement'] = {'mass_g': inputs.get('mass_g'), 'count': inputs.get('count'), 'lot': inputs.get('lot')}
        measured = measurement(record)
    try:
        if measured:
            weight = {'status': 'measured', 'total_g': measured['mass_g'], 'method': 'Báscula · conjunto exacto de la fotografía', 'warning': 'Medición declarada por el usuario.'}
        elif inputs.get('mode') == 'auto':
            weight = estimate(len(grains))
            associated = (payload.get('metadata') or {}).get('rows') or []
            reference_mass = number(associated[0].get('Weight_50K')) if len(associated) == 1 else None
            if reference_mass is not None and reference_mass > 0:
                weight = estimate(len(grains), {'mode':'lot','mass_g':reference_mass,'count':50,'lot':result.get('sample','')})
                weight['method'] = 'Peso de 50 granos del registro asociado'
                weight['warning'] = 'Se aplica el promedio del peso registrado de 50 granos al conteo de esta foto. Es una extrapolación de la muestra asociada, no un pesaje del conjunto fotografiado.'
            if payload.get('review', {}).get('reviewed'):
                try:
                    if reference_mass is None:
                        weight = own_weight_prediction(result)
                except ValueError:
                    pass
        elif inputs.get('mode') == 'trained':
            weight = own_weight_prediction(result) if payload.get('review', {}).get('reviewed') else {'status': 'pending', 'reason': 'Confirma la revisión de granos para usar el modelo de peso.'}
        else:
            weight = estimate(len(grains), inputs)
    except ValueError as exc:
        weight = {'status': 'pending', 'reason': str(exc)}
    yield_result = classify(HISTORY_PATH.parent, payload)
    recommendations = []
    if not payload.get('review', {}).get('reviewed'):
        recommendations.append('En Conteo final, amplía la foto y comprueba un contorno y un número por grano. Corrige objetos extra o segmentos divididos y marca He revisado las detecciones.')
    if result.get('unit') != 'mm':
        recommendations.append('En Zona de análisis y escala, usa dos puntos de la regla original, registra su separación en píxeles y sus milímetros reales; después vuelve a analizar.')
    if any(g.get('toca_borde') for g in grains):
        recommendations.append('Amplía el recorte: hay granos incompletos que tocan el borde.')
    if weight.get('status') == 'scenario':
        recommendations.append('Pesa una referencia del mismo lote para mejorar la aproximación de peso.')
    if yield_result.get('outside_training_range'):
        recommendations.append('Las medidas están fuera del rango del modelo: verifica con una muestra de referencia.')
    if yield_result.get('known_family'):
        recommendations.append('Esta familia participó en el entrenamiento; su clasificación no es una prueba independiente.')
    recommendations.append('Para decidir la aptitud de siembra, complementa estos resultados con pruebas de germinación y vigor.')
    return {'weight': weight, 'yield': yield_result, 'recommendations': recommendations,
            'model_status': 'Activo · experimental' if yield_result.get('model_id') else 'Sin modelo activo'}


@app.post('/api/results/summary')
def summarize_result(payload: dict):
    from .validation import validate_payload
    try:
        return result_summary(validate_payload(payload))
    except ValueError as exc:
        raise HTTPException(422,str(exc))


@app.get("/api/yield/reports/{identifier}.xlsx")
def export_yield_report(identifier: str):
    import json, uuid
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from fastapi.responses import StreamingResponse
    try:
        identifier = str(uuid.UUID(identifier))
    except ValueError:
        raise HTTPException(404, "Reporte no encontrado.")
    source = HISTORY_PATH.parent / "yield_experiments" / (identifier + ".json")
    if not source.exists():
        raise HTTPException(404, "Reporte no encontrado.")
    report = json.loads(source.read_text(encoding="utf-8"))
    book = Workbook()
    book.remove(book.active)
    sheets = {"parametros": [{"parametro": k, "valor": v} for k, v in report.items() if k not in {"samples", "validation", "test", "predictions", "comparison"}],
              "familias": report.get("samples", []), "kernels": report.get("comparison", []),
              "validacion": report.get("validation", []), "prueba": report.get("test", []), "nuevas": report.get("predictions", [])}
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        if rows:
            columns = list(dict.fromkeys(k for row in rows for k in row))
            sheet.append(columns)
            for row in rows:
                values = []
                for key in columns:
                    v = row.get(key)
                    if isinstance(v, (list, dict)):
                        v = json.dumps(v, ensure_ascii=False)
                    if isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@")):
                        v = "'" + v
                    values.append(v)
                sheet.append(values)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                sheet.column_dimensions[cell.column_letter].width = 28
    stream = io.BytesIO()
    book.save(stream)
    stream.seek(0)
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="svm_kernels.xlsx"'})


@app.post("/api/weight/estimate")
def weight_estimate(data: dict):
    from .weight import estimate
    try:
        if (data.get('reference') or {}).get('mode') == 'auto':
            return result_summary({'result': data.get('result') or {}, 'weight_inputs': data['reference'], 'review': {'reviewed': False}})['weight']
        if (data.get("reference") or {}).get("mode") == "trained":
            return own_weight_prediction(data.get("result") or {})
        return estimate(data.get("count"), data.get("reference"))
    except ValueError as exc:
        raise HTTPException(422, str(exc))


def own_weight_prediction(result):
    import json
    from .weight import predict, fingerprint
    directory = HISTORY_PATH.parent / "weight_experiments"
    candidates = sorted(directory.glob("*.json"), key=lambda p: p.stat().st_mtime, reverse=True) if directory.exists() else []
    if not candidates:
        course = HISTORY_PATH.parent/'course_weight.json'
        if course.exists():
            course_report = json.loads(course.read_text(encoding='utf-8'))
            if course_report.get('enabled'):
                prediction = predict(result, course_report)
                prediction['method'] = 'Modelo de peso entrenado con familias del curso'
                prediction['warning'] = 'Validación dejando una familia fuera del curso; error medio para 50 granos: ' + str(round(course_report['mae_total_g'],3)) + ' g. Falta validación con muestras de Coroneo.'
                return prediction
        raise ValueError("Registra pesajes reales para entrenar un modelo de peso local.")
    report = json.loads(candidates[0].read_text(encoding="utf-8"))
    with history_db() as db:
        active = [(i, json.loads(p)) for i, p in db.execute("SELECT id, payload FROM samples WHERE deleted = 0")]
    if report.get("data_fingerprint") != fingerprint(active):
        raise ValueError("El conjunto de entrenamiento cambió. Vuelve a entrenar peso antes de usar el modelo.")
    return predict(result, report)


@app.post("/api/weight/train")
def weight_train():
    import json, uuid
    from datetime import datetime, timezone
    from .weight import training_report
    with history_db() as db:
        rows = db.execute("SELECT id, payload FROM samples WHERE deleted = 0 ORDER BY id").fetchall()
    if not lock.acquire(blocking=False):
        raise HTTPException(429, "Hay un análisis o experimento en curso.")
    try:
        report = training_report([(i, json.loads(p)) for i, p in rows])
        report["created"] = datetime.now(timezone.utc).isoformat()
        if report["status"] == "completed":
            directory = HISTORY_PATH.parent / "weight_experiments"
            directory.mkdir(parents=True, exist_ok=True)
            report["id"] = str(uuid.uuid4())
            path = directory / (report["id"] + ".json")
            temp = path.with_suffix(".tmp")
            temp.write_text(json.dumps(report, ensure_ascii=False, allow_nan=False), encoding="utf-8")
            temp.replace(path)
        return report
    except ImportError:
        raise HTTPException(503, "Instala backend/requirements.txt para entrenar peso.")
    finally:
        lock.release()


@app.get("/api/health")
def health():
    available = importlib.util.find_spec("cellpose") is not None
    return {"status": "ok", "analysis_version": 2, "cellpose_installed": available,
            "methods": ["classical"] + (["cellpose"] if available else [])}


@app.post("/api/analyze")
def analysis(
    file: UploadFile = File(...), sample: str = Form("Muestra"),
    method: str = Form("cellpose"), left: float = Form(0), top: float = Form(0),
    right: float = Form(1), bottom: float = Form(1), diameter: float = Form(0),
    min_area: int = Form(100), max_side: int = Form(1800),
    quality: str = Form("detailed"),
    scale_pixels: float = Form(0), scale_mm: float = Form(0),
):
    import math
    if quality not in {"quick", "detailed"}:
        raise HTTPException(422, "Nivel de análisis no reconocido.")
    if method not in {"cellpose", "classical"}:
        raise HTTPException(422, "Método no reconocido.")
    if not all(math.isfinite(v) for v in [left, top, right, bottom, diameter, scale_pixels, scale_mm]):
        raise HTTPException(422, "Los parámetros deben ser números finitos.")
    if not (0 <= left < right <= 1 and 0 <= top < bottom <= 1):
        raise HTTPException(422, "El recorte debe tener límites entre 0 y 100%, en orden.")
    if not 256 <= max_side <= 2400 or not 1 <= min_area <= 100000 or not 0 <= diameter <= 1000:
        raise HTTPException(422, "Parámetros de segmentación fuera de rango.")
    if scale_pixels < 0 or scale_mm < 0 or bool(scale_pixels) != bool(scale_mm):
        raise HTTPException(422, "Introduce tanto la distancia en píxeles como la distancia real, o deja ambas en cero.")
    if len(sample) > 120:
        raise HTTPException(422, "El nombre de muestra admite hasta 120 caracteres.")
    if method == "cellpose" and importlib.util.find_spec("cellpose") is None:
        raise HTTPException(503, "Cellpose no está instalado en el servidor. Instala backend/requirements.txt o elige explícitamente el método clásico.")
    raw = file.file.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise HTTPException(413, "La imagen supera el límite de 25 MB.")
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(raw)) as source:
                if source.format not in {"JPEG", "PNG", "WEBP"}:
                    raise HTTPException(415, "Utiliza una imagen JPG, PNG o WebP.")
                image = ImageOps.exif_transpose(source).convert("RGB")
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError, Image.DecompressionBombWarning):
        raise HTTPException(422, "No se pudo abrir la imagen o supera el límite de 85 megapíxeles.")
    if not lock.acquire(blocking=False):
        raise HTTPException(429, "Hay un análisis en curso. Espera a que termine.")
    try:
        started_analysis = time.perf_counter()
        target_side = min(max_side, 1000) if quality == "quick" else max_side
        width, height = image.size
        crop_side = max(int(right * width) - int(left * width), int(bottom * height) - int(top * height))
        reference_factor = min(1.0, max_side / max(1, crop_side))
        working_factor = min(1.0, target_side / max(1, crop_side))
        ratio = working_factor / reference_factor
        adjusted_area = max(1, round(min_area * ratio ** 2))
        adjusted_diameter = diameter * ratio if diameter else None
        result = analyze(image, method=method, roi=(left, top, right, bottom), diameter=adjusted_diameter,
                       min_area=adjusted_area, max_side=target_side, scale_pixels=scale_pixels or None,
                       scale_mm=scale_mm or None, sample=sample.strip() or "Muestra")
        result["quality"] = quality
        result["effective_min_area"] = adjusted_area
        result["effective_diameter"] = adjusted_diameter
        if quality == "quick":
            result["warnings"].append("Modo rápido: la resolución reducida puede perder detalles de contorno o granos pequeños. Usa Detallado para revisar las medidas.")
        result["processing_seconds"] = round(time.perf_counter() - started_analysis, 3)
        logging.getLogger("uvicorn.error").info("Análisis %s: %s granos en %.3f s", method, result["count"], result["processing_seconds"])
        return result
    except ValueError as exc:
        raise HTTPException(422, str(exc))
    except Exception:
        logging.exception("Error durante el análisis")
        raise HTTPException(500, "El análisis falló. Consulta el registro del servidor; si usas Cellpose, revisa la instalación y la descarga de pesos.")
    finally:
        lock.release()


@app.post("/api/metadata")
def metadata(file: UploadFile = File(...)):
    import csv
    name = (file.filename or "").lower()
    raw = file.file.read(5 * 1024 * 1024 + 1)
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(413, "Los metadatos superan el límite de 5 MB.")
    try:
        if name.endswith(".csv"):
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("latin-1")
            if text.startswith("sep="):
                separator, text = text[4:5], text.split("\n", 1)[1]
            else:
                try:
                    separator = csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
                except csv.Error:
                    separator = ","
            records = csv.reader(io.StringIO(text), delimiter=separator)
        elif name.endswith(".xlsx"):
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            records = book.active.iter_rows(values_only=True)
        else:
            raise HTTPException(415, "Utiliza un archivo CSV o Excel .xlsx.")
        columns = [str(v).strip() if v is not None else "" for v in next(records)]
        if not columns or any(not c for c in columns) or len(set(columns)) != len(columns):
            raise ValueError("La primera fila debe contener encabezados únicos y no vacíos.")
        if len(columns) > 100:
            raise ValueError("El archivo admite hasta 100 columnas.")
        rows = []
        for values in records:
            if not any(v is not None and str(v).strip() for v in values):
                continue
            if len(values) != len(columns):
                raise ValueError("Hay filas con una cantidad de columnas distinta a los encabezados.")
            rows.append({c: str(v) if v is not None else "" for c, v in zip(columns, values)})
            if len(rows) > 5000:
                raise ValueError("El archivo admite hasta 5000 muestras.")
        if not rows:
            raise ValueError("El archivo no contiene muestras.")
        return {"columns": columns, "rows": rows}
    except HTTPException:
        raise
    except ImportError:
        raise HTTPException(503, "Instala las dependencias actualizadas del backend para leer Excel.")
    except Exception as exc:
        detail = str(exc) if isinstance(exc, ValueError) else "No se pudo leer el archivo. Revisa su formato y encabezados."
        raise HTTPException(422, detail)
    finally:
        if "book" in locals():
            book.close()


METADATA_PATH = __import__("pathlib").Path(__file__).resolve().parent.parent / "data" / "metadata.json"
metadata_lock = threading.Lock()

@app.get("/api/metadata/saved")
def saved_metadata():
    import json
    with metadata_lock:
        if not METADATA_PATH.exists():
            return {"available": False}
        data = json.loads(METADATA_PATH.read_text(encoding="utf-8"))
    return {"available": True, **data}


@app.post('/api/metadata/record')
def register_local_metadata(record: dict):
    import json
    from .yield_experiment import number, family
    identifier = str(record.get('nombre_nuevo','')).strip()
    if not identifier or len(identifier)>120:
        raise HTTPException(422,'Escribe un identificador de muestra de 1 a 120 caracteres.')
    allowed = ['nombre_nuevo','localidad','campana','lote','Weight_50K','Yield','unidad_rendimiento']
    row = {key:str(record.get(key,'')).strip() for key in allowed}
    for key in ['Weight_50K','Yield']:
        if row[key] and (number(row[key]) is None or number(row[key])<0 or (key=='Weight_50K' and number(row[key])==0)):
            raise HTTPException(422,'Los pesos y el rendimiento deben ser números reales válidos.')
    row['procedencia']='Registro local del usuario'
    if row['Yield'] and row['unidad_rendimiento'] not in ['t/ha','kg/ha']:
        raise HTTPException(422,'Indica la unidad del rendimiento medido.')
    if row['Yield'] and row['unidad_rendimiento']=='kg/ha':
        row['Yield']=str(float(row['Yield'])/1000)
        row['unidad_rendimiento']='t/ha'
    with metadata_lock:
        data = json.loads(METADATA_PATH.read_text(encoding='utf-8')) if METADATA_PATH.exists() else {'filename':'muestras_SIMAC','columns':[],'rows':[]}
        identifier_key=next((key for key in data['columns'] if key in ['nombre_nuevo','nombre_imagen','familia','muestra']),'nombre_nuevo')
        row[identifier_key]=identifier
        for saved in data['rows']:
            if any(family(saved.get(k,''))==family(identifier) for k in ['nombre_nuevo','nombre_imagen','familia']):
                raise HTTPException(409,'Ese identificador ya está registrado. Usa un identificador nuevo para la muestra local, sin reemplazar datos del curso.')
        for key in row:
            if key not in data['columns']: data['columns'].append(key)
        data['rows'].append(row)
        METADATA_PATH.parent.mkdir(parents=True,exist_ok=True)
        temp=METADATA_PATH.with_suffix('.tmp');temp.write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8');temp.replace(METADATA_PATH)
    return {'available':True,**data}

@app.post("/api/metadata/saved")
def save_metadata(file: UploadFile = File(...)):
    import json
    from pathlib import Path
    data = {**metadata(file), "filename": Path(file.filename or "metadatos").name}
    with metadata_lock:
        METADATA_PATH.parent.mkdir(parents=True, exist_ok=True)
        temporary = METADATA_PATH.with_suffix(".tmp")
        temporary.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        temporary.replace(METADATA_PATH)
    return {"available": True, **data}


HISTORY_PATH = METADATA_PATH.parent / "samples.sqlite3"

from contextlib import contextmanager

@contextmanager
def history_db():
    import sqlite3
    HISTORY_PATH.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(HISTORY_PATH)
    connection.execute("CREATE TABLE IF NOT EXISTS samples (id TEXT PRIMARY KEY, created TEXT, name TEXT, filename TEXT, image BLOB, payload TEXT)")
    if "deleted" not in {row[1] for row in connection.execute("PRAGMA table_info(samples)")}:
        connection.execute("ALTER TABLE samples ADD COLUMN deleted INTEGER NOT NULL DEFAULT 0")
    try:
        with connection:
            yield connection
    finally:
        connection.close()

@app.post("/api/samples")
def save_sample(file: UploadFile = File(...), payload: UploadFile = File(...), identifier: str = Form("")):
    import json, uuid
    from datetime import datetime, timezone
    payload_raw = payload.file.read(40 * 1024 * 1024 + 1)
    if len(payload_raw) > 40 * 1024 * 1024:
        raise HTTPException(413, "El resultado supera el límite de guardado.")
    if payload_raw.startswith(b'\x1f\x8b'):
        import gzip, io
        try:
            with gzip.GzipFile(fileobj=io.BytesIO(payload_raw)) as compressed:
                payload_raw = compressed.read(160 * 1024 * 1024 + 1)
        except (OSError, EOFError):
            raise HTTPException(422, "El archivo de resultados comprimido no es válido.")
        if len(payload_raw) > 160 * 1024 * 1024:
            raise HTTPException(413, "El análisis contiene demasiadas vistas para guardarlo. Guarda antes de realizar más uniones.")
    try:
        payload_text = payload_raw.decode("utf-8")
        data = json.loads(payload_text)
        from .validation import validate_payload
        validate_payload(data)
        result = data["result"]
        if not isinstance(result, dict) or not isinstance(result.get("sample"), str) or not isinstance(result.get("grains"), list):
            raise ValueError()
    except (ValueError, KeyError, TypeError, UnicodeDecodeError):
        raise HTTPException(422, "El resultado de la muestra no es válido.")
    from .weight import estimate, measurement
    inputs = data.get("weight_inputs") or {"mode": "literature"}
    try:
        data["weight_estimate"] = result_summary(data)['weight'] if inputs.get('mode') == 'auto' else own_weight_prediction(result) if inputs.get("mode") == "trained" else estimate(len(result["grains"]), inputs)
    except ValueError as exc:
        raise HTTPException(422, str(exc))
    data.pop("weight_measurement", None)
    if inputs.get("measured"):
        signature = [[g.get(k) for k in ["grano", "largo", "ancho", "area"]] for g in result["grains"]]
        if inputs.get("grain_signature") != signature or inputs.get("mode") != "lot":
            raise HTTPException(422, "Los granos cambiaron desde el pesaje. Confirma un nuevo pesaje del conjunto exacto.")
        data["weight_measurement"] = {"mass_g": inputs.get("mass_g"), "count": inputs.get("count"), "lot": inputs.get("lot"), "moisture_pct": inputs.get("moisture_pct")}
        valid = measurement(data)
        if valid is None:
            raise HTTPException(422, "El peso real requiere revisión, conteo exacto y lote.")
        data["weight_measurement"] = valid
    data['application_summary'] = result_summary(data)
    payload_text = json.dumps(data, ensure_ascii=False, allow_nan=False)
    raw = file.file.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise HTTPException(413, "La fotografía supera 25 MB.")
    existing_id = identifier
    identifier = identifier or str(uuid.uuid4())
    created = datetime.now(timezone.utc).isoformat()
    with history_db() as db:
        if existing_id:
            existing = db.execute("SELECT created FROM samples WHERE id = ? AND deleted = 0", (identifier,)).fetchone()
            if existing is None:
                raise HTTPException(404, "No se encontró la muestra para actualizar.")
            created = existing[0]
            db.execute("UPDATE samples SET name = ?, filename = ?, image = ?, payload = ? WHERE id = ?", (result["sample"], file.filename, raw, payload_text, identifier))
        else:
            db.execute("INSERT INTO samples (id, created, name, filename, image, payload) VALUES (?, ?, ?, ?, ?, ?)", (identifier, created, result["sample"], file.filename, raw, payload_text))
    return {"id": identifier, "created": created}

@app.get("/api/samples")
def list_samples():
    import json
    with history_db() as db:
        records = db.execute("SELECT id, created, name, payload FROM samples WHERE deleted = 0 ORDER BY created DESC").fetchall()
    return [{"id": i, "created": date, "name": name, "count": json.loads(payload)["result"].get("count", 0)} for i, date, name, payload in records]

@app.get("/api/samples/{identifier}")
def get_sample(identifier: str):
    import json, base64
    with history_db() as db:
        row = db.execute("SELECT created, filename, image, payload FROM samples WHERE id = ? AND deleted = 0", (identifier,)).fetchone()
    if row is None:
        raise HTTPException(404, "No se encontró la muestra.")
    created, filename, raw, payload = row
    return {"id": identifier, "created": created, "filename": filename, "image_base64": base64.b64encode(raw).decode(), **json.loads(payload)}


@app.post("/api/samples/merge")
def merge_detections(file: UploadFile = File(...), payload: UploadFile = File(...)):
    import json, base64
    import numpy as np
    raw_payload = payload.file.read(40 * 1024 * 1024 + 1)
    raw = file.file.read(MAX_BYTES + 1)
    if len(raw_payload) > 40 * 1024 * 1024 or len(raw) > MAX_BYTES:
        raise HTTPException(413, "La muestra supera el límite permitido.")
    try:
        data = json.loads(raw_payload)
        source, parameters = data["result"], data["parameters"]
        pair = data["pair"]
        ids = {g["grano"] for g in source["grains"]}
        if len(pair) != 2 or pair[0] == pair[1] or not all(i in ids for i in pair):
            raise ValueError("Selecciona dos detecciones diferentes.")
        encoded = source["editable_masks"].split(",", 1)[1]
        with Image.open(io.BytesIO(base64.b64decode(encoded, validate=True))) as mask_image:
            if max(mask_image.size) > 2400:
                raise ValueError("Máscara fuera de rango.")
            masks = np.asarray(mask_image).astype(np.int32).copy()
        # La unión admite segmentos cuyos límites se superponen o se tocan.
        a, b = masks == pair[0], masks == pair[1]
        from skimage.morphology import dilation, disk
        if not np.any(dilation(a, footprint=disk(8)) & b):
            raise ValueError("Las detecciones están separadas. Une solamente partes del mismo grano.")
        masks[b] = pair[0]
        with Image.open(io.BytesIO(raw)) as photo:
            image = ImageOps.exif_transpose(photo).convert("RGB")
        roi = tuple(float(v) / 100 for v in parameters["roi"])
        result = analyze(image, method=parameters["method"], roi=roi, diameter=None, min_area=1,
            max_side=max(source["working_size"]), scale_pixels=float(parameters["scalePx"] or 0) or None,
            scale_mm=float(parameters["scaleMm"] or 0) or None, sample=source["sample"], supplied_masks=masks)
        result["quality"] = source.get("quality", "detailed")
        result["warnings"].append("Se unieron dos detecciones por revisión manual. Comprueba que correspondan al mismo grano.")
        return result
    except (ValueError, KeyError, TypeError, OSError, IndexError) as exc:
        raise HTTPException(422, str(exc) if isinstance(exc, ValueError) else "No se pudo unir. Realiza un nuevo análisis para obtener máscaras editables.")


@app.patch("/api/samples/{identifier}")
def rename_sample(identifier: str, name: str = Form(...)):
    import json, csv
    name = name.strip()
    if not name or len(name) > 120:
        raise HTTPException(422, "Escribe un nombre de 1 a 120 caracteres.")
    with history_db() as db:
        row = db.execute("SELECT payload FROM samples WHERE id = ? AND deleted = 0", (identifier,)).fetchone()
        if row is None:
            raise HTTPException(404, "No se encontró la muestra.")
        data = json.loads(row[0])
        def rename_result(result):
            if not result:
                return
            result["sample"] = name
            for grain in result.get("grains", []):
                grain["muestra"] = name
            if result.get("csv"):
                reader = csv.DictReader(io.StringIO(result["csv"]))
                output = io.StringIO(newline="")
                writer = csv.DictWriter(output, fieldnames=reader.fieldnames)
                writer.writeheader()
                safe = "'" + name if name.lstrip().startswith(("=", "+", "-", "@")) else name
                for record in reader:
                    if "muestra" in record:
                        record["muestra"] = safe
                    writer.writerow(record)
                result["csv"] = output.getvalue()
        rename_result(data.get("result"))
        rename_result(data.get("original_result"))
        for result in data.get("merge_history", []):
            rename_result(result)
        db.execute("UPDATE samples SET name = ?, payload = ? WHERE id = ?", (name, json.dumps(data, ensure_ascii=False), identifier))
    return {"id": identifier, "name": name}

@app.delete("/api/samples/{identifier}")
def delete_sample(identifier: str):
    with history_db() as db:
        cursor = db.execute("UPDATE samples SET deleted = 1 WHERE id = ? AND deleted = 0", (identifier,))
        if not cursor.rowcount:
            raise HTTPException(404, "No se encontró la muestra.")
    return {"id": identifier, "deleted": True}

@app.post("/api/samples/{identifier}/restore")
def restore_sample(identifier: str):
    with history_db() as db:
        cursor = db.execute("UPDATE samples SET deleted = 0 WHERE id = ? AND deleted = 1", (identifier,))
        if not cursor.rowcount:
            raise HTTPException(404, "No se encontró la muestra eliminada.")
    return {"id": identifier, "deleted": False}


from .agriculture import router as agriculture_router
app.include_router(agriculture_router(lambda: HISTORY_PATH))

from .assistant import router as assistant_router
app.include_router(assistant_router)
