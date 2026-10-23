from backend.tests.support import TEMP_ROOT
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from fastapi import UploadFile, HTTPException
from backend.weight import estimate, measurement, training_report, predict, fingerprint
from backend.tests.test_yield_experiment import sample


class WeightTests(unittest.TestCase):
    def test_literature_is_scenario_not_measured(self):
        result = estimate(50)
        self.assertEqual(result["status"], "scenario")
        self.assertAlmostEqual(result["total_range_g"][0], 14.175, places=3)
        self.assertNotIn("total_g", result)
        self.assertEqual(estimate(0)["status"], "unavailable")

    def test_lot_math_and_invalid_reference(self):
        result = estimate(100, {"mode": "lot", "mass_g": 15, "count": 50})
        self.assertEqual(result["total_g"], 30)
        for mass, count in [(float("nan"), 50), (10, 0), (10, 2.5)]:
            with self.assertRaises(ValueError): estimate(50, {"mode": "lot", "mass_g": mass, "count": count})

    def test_only_exact_measurements_train(self):
        data = sample()
        data["weight_estimate"] = estimate(25)
        self.assertIsNone(measurement(data))
        data["weight_measurement"] = {"mass_g": 7, "count": 24, "lot": "A"}
        self.assertIsNone(measurement(data))
        data["weight_measurement"]["count"] = 25
        self.assertEqual(measurement(data)["mass_g"], 7)
        self.assertEqual(training_report([("1", data)])["status"], "blocked")

    def test_training_group_validation_and_prediction(self):
        records = []
        for i in range(12):
            data = sample(i)
            data["weight_measurement"] = {"mass_g": 5+i*.2, "count": 25, "lot": f"L{i%3}"}
            records.append((str(i), data))
        report = training_report(records)
        self.assertEqual(report["status"], "completed")
        self.assertEqual(len(report["validation"]), 12)
        report["id"] = "test-model"
        result = predict(sample(3)["result"], report)
        self.assertGreater(result["total_g"], 0)
        self.assertEqual(result["model_id"], "test-model")
        prior = fingerprint(records)
        records[0][1]["weight_measurement"]["mass_g"] = 9
        self.assertNotEqual(prior, fingerprint(records))

    def test_save_updates_same_sample_and_keeps_weight_provenance(self):
        from backend import main
        data = sample()
        data["weight_inputs"] = {"mode": "lot", "mass_g": "7", "count": "25", "lot": "L1", "measured": True,
            "grain_signature": [[g.get(k) for k in ["grano", "largo", "ancho", "area"]] for g in data["result"]["grains"]]}
        def save(identifier=""):
            return main.save_sample(UploadFile(filename="test.jpg", file=io.BytesIO(b"photo")),
                UploadFile(filename="payload.json", file=io.BytesIO(json.dumps(data).encode())), identifier=identifier)
        with tempfile.TemporaryDirectory(dir=TEMP_ROOT) as folder, patch.object(main, "HISTORY_PATH", Path(folder)/"samples.sqlite3"):
            saved = save()
            self.assertEqual(save(saved["id"])["id"], saved["id"])
            self.assertEqual(len(main.list_samples()), 1)
            loaded = main.get_sample(saved["id"])
            self.assertEqual(loaded["weight_measurement"]["mass_g"], 7)
            self.assertEqual(loaded["weight_estimate"]["total_g"], 7)
            data["weight_inputs"]["grain_signature"] = []
            with self.assertRaises(HTTPException): save(saved["id"])

    def test_excel_report_export_and_formula_escape(self):
        import asyncio, uuid
        from openpyxl import load_workbook
        from backend import main
        identifier = str(uuid.uuid4())
        report = {"id": identifier, "status": "completed", "samples": [{"family": "=1+1", "ready": True}],
                  "parameters": {"kernel": "rbf"}, "test": [], "validation": [], "predictions": []}
        with tempfile.TemporaryDirectory(dir=TEMP_ROOT) as folder, patch.object(main, "HISTORY_PATH", Path(folder)/"samples.sqlite3"):
            directory = Path(folder)/"yield_experiments"; directory.mkdir()
            (directory/(identifier+".json")).write_text(json.dumps(report), encoding="utf-8")
            response = main.export_yield_report(identifier)
            async def read():
                return b"".join([chunk async for chunk in response.body_iterator])
            book = load_workbook(io.BytesIO(asyncio.run(read())))
            self.assertEqual(book["familias"]["A2"].value, "'=1+1")
            self.assertIn("parametros", book.sheetnames)
            book.close()
